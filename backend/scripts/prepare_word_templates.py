"""
Обновление шаблонов в backend/templates/ (только для разработки).

На сервере используются уже готовые файлы:
  backend/templates/addendum_template.docx
  backend/templates/framework_template.docx

Положите исходники в backend/templates/_sources/ и запустите:
  python -m scripts.prepare_word_templates
"""
import os
import shutil
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from docx import Document

from app.core.config import BACKEND_DIR, get_templates_dir

SOURCES = get_templates_dir() / "_sources"
DEFAULT_ADDENDUM = SOURCES / "addendum_source.docx"
DEFAULT_FRAMEWORK = SOURCES / "framework_source.docx"


def _replace_in_paragraphs(doc: Document, mapping: dict[str, str]) -> None:
    for para in doc.paragraphs:
        text = para.text
        if not text:
            continue
        new = text
        for old, repl in mapping.items():
            if old in new:
                new = new.replace(old, repl)
        if new != text:
            para.text = new
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    text = para.text
                    if not text:
                        continue
                    new = text
                    for old, repl in mapping.items():
                        if old in new:
                            new = new.replace(old, repl)
                    if new != text:
                        para.text = new


def _find_product_table(doc: Document):
    for table in doc.tables:
        if len(table.columns) != 3:
            continue
        header = [c.text.strip() for c in table.rows[0].cells]
        if "Наименование документа" in header[0] and "Продукция" in header[1]:
            return table
    return None


def prepare_addendum(source: Path, dest: Path) -> None:
    shutil.copy2(source, dest)
    doc = Document(str(dest))

    mapping = {
        "Приложение №3": "Приложение №{{ addendum_number }}",
        "к договору № ФВ-025-2020 от 22 апреля 2020 года": (
            "к договору № {{ framework_number }} от {{ framework_date_text }}"
        ),
        "Приложение №3 к договору № ФВ-025-2020 от 22 апреля 2020 года": (
            "Приложение №{{ addendum_number }} к договору № {{ framework_number }} "
            "от {{ framework_date_text }}"
        ),
        "30 (тридцать) рабочих дней": "{{ work_days_words }}",
        "125 000 руб. 00 коп.": "{{ total_amount_formatted }}",
        "«19» декабря 2022 г.": "«{{ addendum_date_day }}» {{ addendum_date_month }} {{ addendum_date_year }} г.",
    }
    _replace_in_paragraphs(doc, mapping)

    for para in doc.paragraphs:
        if "именуемое в дальнейшем «Заказчик»" in para.text:
            para.text = "{{ customer_preamble }}"
            break

    product_table = _find_product_table(doc)
    if product_table:
        while len(product_table.rows) > 2:
            product_table._tbl.remove(product_table.rows[-1]._tr)
        for cell in product_table.rows[1].cells:
            cell.text = ""

    for table in doc.tables:
        if len(table.columns) != 2:
            continue
        if table.rows[0].cells[0].text.strip().upper() != "ИСПОЛНИТЕЛЬ":
            continue
        if len(table.rows) < 2:
            continue
        table.rows[1].cells[0].paragraphs[0].text = "{{ executor_short_name }}"
        table.rows[1].cells[1].paragraphs[0].text = "{{ customer_short_name }}"
        break

    doc.save(str(dest))
    print(f"  -> {dest.name}")


def prepare_framework(source: Path, dest: Path) -> None:
    shutil.copy2(source, dest)
    doc = Document(str(dest))

    for para in doc.paragraphs:
        if "Договор №" in para.text and "{{" not in para.text:
            para.text = "Договор № {{ contract_number }}"
        if "именуемое в дальнейшем «Заказчик»" in para.text:
            para.text = "{{ customer_preamble }}"
            break
        if "именуемое в дальнейшем «Исполнитель»" in para.text:
            para.text = "{{ executor_preamble }}"
            break

    for table in doc.tables:
        if len(table.columns) == 2 and table.rows[0].cells[0].text.strip().upper() == "ИСПОЛНИТЕЛЬ":
            table.rows[1].cells[0].paragraphs[0].text = "{{ executor_short_name }}"
            table.rows[1].cells[1].paragraphs[0].text = "{{ customer_short_name }}"
            break

    doc.save(str(dest))
    print(f"  -> {dest.name}")


def main():
    out_dir = get_templates_dir()
    out_dir.mkdir(parents=True, exist_ok=True)
    SOURCES.mkdir(parents=True, exist_ok=True)

    add_src = Path(os.environ.get("ADDENDUM_SOURCE", DEFAULT_ADDENDUM))
    fw_src = Path(os.environ.get("FRAMEWORK_SOURCE", DEFAULT_FRAMEWORK))

    print(f"Шаблоны: {out_dir}")
    if not add_src.exists():
        print(f"Нет исходника приложения: {add_src}")
        print("Скопируйте ваш .docx в templates/_sources/addendum_source.docx")
        sys.exit(1)
    if not fw_src.exists():
        print(f"Нет исходника рамки: {fw_src}")
        print("Скопируйте ваш .docx в templates/_sources/framework_source.docx")
        sys.exit(1)

    prepare_addendum(add_src, out_dir / "addendum_template.docx")
    prepare_framework(fw_src, out_dir / "framework_template.docx")
    print("Готово. Закоммитьте backend/templates/*.docx в репозиторий.")


if __name__ == "__main__":
    main()
