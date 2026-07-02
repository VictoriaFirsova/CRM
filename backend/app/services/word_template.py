"""Генерация .docx из шаблонов (оформление как в образцах NSK)."""
from __future__ import annotations

from copy import deepcopy
from datetime import date
from pathlib import Path

from docx import Document
from docxtpl import DocxTemplate
from docx.oxml.ns import qn
from docx.shared import Pt

from app.core.config import get_templates_dir
from app.models.client import Client
from app.models.contract import Contract
from app.models.enums import ContractKind
from app.services.executor import _format_line_price, _format_rub, build_document_variables

MONTHS_GENITIVE = (
    "января",
    "февраля",
    "марта",
    "апреля",
    "мая",
    "июня",
    "июля",
    "августа",
    "сентября",
    "октября",
    "ноября",
    "декабря",
)


def _date_text(d: date | None) -> str:
    if not d:
        return ""
    return f"{d.day} {MONTHS_GENITIVE[d.month - 1]} {d.year} года"


def _work_days_words(n: int | None) -> str:
    if not n:
        return ""
    return f"{n} ({_number_words_simple(n)}) рабочих дней"


def _number_words_simple(n: int) -> str:
    words = {30: "тридцать", 45: "сорок пять", 60: "шестьдесят", 90: "девяносто"}
    return words.get(n, str(n))


def template_path(kind: ContractKind) -> Path:
    name = "addendum_template.docx" if kind == ContractKind.ADDENDUM else "framework_template.docx"
    return get_templates_dir() / name


def templates_available() -> bool:
    return template_path(ContractKind.ADDENDUM).is_file() and template_path(
        ContractKind.FRAMEWORK
    ).is_file()


def build_template_context(contract: Contract, client: Client | None) -> dict:
    variables = build_document_variables(contract, client)
    ctx: dict = dict(variables)

    addendum_date = contract.start_date or date.today()
    framework_date = None
    if contract.parent and contract.parent.start_date:
        framework_date = contract.parent.start_date

    ctx.update(
        {
            "addendum_date_day": str(addendum_date.day),
            "addendum_date_month": MONTHS_GENITIVE[addendum_date.month - 1],
            "addendum_date_year": str(addendum_date.year),
            "framework_date_text": _date_text(framework_date),
            "work_days_words": _work_days_words(contract.work_days),
            "total_amount_formatted": variables.get("total_amount", ""),
            "advance_percent": str(contract.advance_percent or 100),
            "customer_short_name": variables.get("customer_short_name", ""),
            "executor_short_name": variables.get("executor_short_name", ""),
        }
    )
    return ctx


def _find_product_table(doc: Document):
    """Таблица №1: 3 колонки, заголовок с документом/продукцией/стоимостью."""
    for table in doc.tables:
        if len(table.columns) != 3 or not table.rows:
            continue
        header = [c.text.strip().lower() for c in table.rows[0].cells]
        h0, h1 = header[0], header[1] if len(header) > 1 else ""
        if "исполнитель" in h0:
            continue
        if (
            "наименование документа" in h0
            or ("документ" in h0 and "продукц" in h1)
            or ("документ" in h0 and ("стоим" in header[2] if len(header) > 2 else False))
        ):
            return table
    return None


def _fill_product_table(doc: Document, contract: Contract) -> None:
    table = _find_product_table(doc)
    if not table or len(table.rows) < 2:
        return

    lines = sorted(contract.line_items or [], key=lambda x: x.position)
    template_tr = table.rows[1]._tr

    while len(table.rows) > 2:
        table._tbl.remove(table.rows[-1]._tr)

    if not lines:
        # Одна пустая строка по умолчанию
        for cell in table.rows[1].cells:
            _set_cell_text_with_font(cell, "")
        return

    for idx, line in enumerate(lines):
        if idx == 0:
            row = table.rows[1]
        else:
            new_tr = deepcopy(template_tr)
            table._tbl.append(new_tr)
            row = table.rows[-1]
        _set_cell_text_with_font(row.cells[0], line.document_name)
        _set_cell_text_with_font(row.cells[1], line.product_name)
        _set_cell_text_with_font(row.cells[2], _format_line_price(line.price, line.vat_exempt))


def _set_cell_text_with_font(cell, text: str) -> None:
    """Записывает текст в ячейку и фиксирует шрифт таблицы."""
    cell.text = text or ""
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(10)
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")


def render_from_template(contract: Contract, client: Client | None, output_path: Path) -> bool:
    tpl_path = template_path(contract.kind)
    if not tpl_path.is_file():
        return False

    ctx = build_template_context(contract, client)
    tmp = output_path.with_suffix(".tmp.docx")
    doc_tpl = DocxTemplate(str(tpl_path))
    doc_tpl.render(ctx)
    doc_tpl.save(str(tmp))

    doc = Document(str(tmp))
    if contract.kind == ContractKind.ADDENDUM:
        _fill_product_table(doc, contract)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    tmp.unlink(missing_ok=True)
    return True
