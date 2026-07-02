import re
from io import BytesIO

from docx import Document


def _read_docx_text(file_bytes: bytes) -> str:
    doc = Document(BytesIO(file_bytes))
    lines: list[str] = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.replace("\xa0", " ").strip()
        if text:
            lines.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.replace("\xa0", " ").strip() for c in row.cells]
            row_text = " | ".join(c for c in cells if c)
            if row_text:
                lines.append(row_text)
    return "\n".join(lines)


def _trim_at_pattern(value: str, pattern: str | None) -> str:
    """Обрезает значение до первого вхождения pattern (например, до «БИК»)."""
    if not value or not pattern:
        return value.strip(" .,;") if value else ""
    m = re.search(pattern, value, re.IGNORECASE)
    if m:
        value = value[: m.start()]
    return value.strip(" .,;")


def _search_label(
    text: str,
    labels: list[str],
    stop_before: str | None = None,
) -> str | None:
    """Ищет значение после метки в той же строке или на следующей."""
    lines = text.splitlines()
    for i, line in enumerate(lines):
        for label in labels:
            if label.lower() not in line.lower():
                continue
            parts = re.split(rf"{re.escape(label)}\s*[:№]?\s*", line, maxsplit=1, flags=re.IGNORECASE)
            if len(parts) > 1 and parts[1].strip():
                value = _trim_at_pattern(parts[1], stop_before)
                if value:
                    return value
            if i + 1 < len(lines) and lines[i + 1].strip():
                value = _trim_at_pattern(lines[i + 1], stop_before)
                if value:
                    return value
    return None


def _first_match(text: str, pattern: str) -> str | None:
    m = re.search(pattern, text, re.IGNORECASE | re.MULTILINE)
    if not m:
        return None
    try:
        return m.group(1).strip()
    except IndexError:
        return m.group(0).strip()


def _all_matches(text: str, pattern: str) -> list[str]:
    return [m.group(1) for m in re.finditer(pattern, text, re.IGNORECASE | re.MULTILINE)]


def parse_client_requisites(file_bytes: bytes) -> dict:
    """
    Извлекает реквизиты из DOCX. Все поля опциональны — пользователь может поправить вручную.
    """
    text = _read_docx_text(file_bytes)
    if not text.strip():
        return {"fields": {}, "filled": [], "warnings": ["Файл пуст или не содержит текста"]}

    fields: dict[str, str | None] = {}
    filled: list[str] = []
    warnings: list[str] = []

    inn = _first_match(text, r"(?:ИНН|инн)\s*[:№]?\s*(\d{10}|\d{12})\b")
    if not inn:
        inns = _all_matches(text, r"\b(\d{10}|\d{12})\b")
        inn = inns[0] if inns else None
    fields["inn"] = inn

    fields["bik"] = _first_match(text, r"(?:БИК|бик)\s*[:№]?\s*(\d{9})\b")

    fields["kpp"] = _first_match(text, r"(?:КПП|кпп)\s*[:№]?\s*(\d{9})\b")
    if not fields["kpp"]:
        bik_val = fields["bik"]
        kpps = [k for k in _all_matches(text, r"\b(\d{9})\b") if k != bik_val]
        fields["kpp"] = kpps[0] if kpps else None

    fields["ogrn"] = _first_match(text, r"(?:ОГРНИП|ОГРН|огрн(?:ип)?)\s*[:№]?\s*(\d{13}|\d{15})\b")

    fields["settlement_account"] = _first_match(
        text,
        r"(?:р\.?\s*/?\s*с\.?|расч[её]тный\s+сч[её]т)\s*[:№]?\s*(\d{20})\b",
    )
    fields["correspondent_account"] = _first_match(
        text,
        r"(?:к\.?\s*/?\s*с\.?|корр\.?\s*сч[её]т|корреспондентский\s+сч[её]т)\s*[:№]?\s*(\d{20})\b",
    )

    accounts_20 = _all_matches(text, r"\b(\d{20})\b")
    if accounts_20:
        if not fields["settlement_account"]:
            fields["settlement_account"] = accounts_20[0]
        if not fields["correspondent_account"] and len(accounts_20) > 1:
            fields["correspondent_account"] = accounts_20[1]
        elif not fields["correspondent_account"] and len(accounts_20) == 1:
            pass

    # Название банка и к/с — из DaData по БИК (см. enrich_client_bank_fields)
    fields["bank_name"] = None

    fields["legal_address"] = _search_label(
        text,
        ["Юридический адрес", "Юр. адрес", "юр. адрес", "Адрес (юридический)"],
    )
    fields["actual_address"] = _search_label(
        text,
        ["Фактический адрес", "Факт. адрес", "Почтовый адрес", "Адрес (фактический)"],
    )

    fields["director_name"] = _search_label(
        text,
        ["Генеральный директор", "Директор", "Руководитель", "В лице"],
    )
    if fields["director_name"]:
        fields["director_name"] = re.sub(
            r"^(генеральный\s+директор|директор|руководитель)\s*",
            "",
            fields["director_name"],
            flags=re.IGNORECASE,
        ).strip()
    fields["director_name_genitive"] = fields["director_name"]
    fields["signer_position"] = _search_label(
        text,
        ["Должность", "Подписант", "Руководитель", "В лице"],
    )
    fields["signer_basis"] = _search_label(
        text,
        ["Действует на основании", "Основание полномочий", "Основание"],
    )

    fields["contact_person"] = _search_label(text, ["Контактное лицо", "Контакт"])
    fields["phone"] = _first_match(
        text,
        r"(?:тел\.?|телефон)\s*[:№]?\s*([+\d\s()\-]{7,20})",
    )
    fields["email"] = _first_match(
        text,
        r"([a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,})",
    )

    fields["full_company_name"] = _search_label(
        text,
        ["Полное наименование", "Наименование организации", "Организация"],
    )
    fields["company_name"] = fields["full_company_name"]

    if not fields["company_name"]:
        for line in text.splitlines()[:15]:
            if re.search(r"\b(ООО|ОАО|ПАО|АО|ИП|ЗАО)\b", line, re.IGNORECASE):
                fields["company_name"] = line.strip()
                if len(line) > 40:
                    fields["full_company_name"] = line.strip()
                break

    if not fields["company_name"]:
        first_line = next((ln.strip() for ln in text.splitlines() if len(ln.strip()) > 3), None)
        fields["company_name"] = first_line

    for key, value in list(fields.items()):
        if value:
            value = value.strip()
            fields[key] = value
            filled.append(key)
        else:
            fields[key] = None

    if not filled:
        warnings.append("Не удалось распознать реквизиты. Проверьте формат файла или заполните вручную.")
    elif len(filled) < 3:
        warnings.append("Распознано мало полей — обязательно проверьте данные перед сохранением.")

    return {"fields": fields, "filled": filled, "warnings": warnings}
