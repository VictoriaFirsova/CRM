import os
import shutil
from datetime import datetime, timezone
from pathlib import Path
from uuid import uuid4

from fastapi import HTTPException, UploadFile, status

from app.core.config import get_templates_dir, settings
from app.core.permissions import can_edit_contract
from app.models.document import Document
from app.models.enums import DocumentType
from app.models.user import User
from app.repositories.contract import ContractRepository
from app.repositories.document import DocumentRepository
from app.schemas.document import DocumentGenerateRequest, DocumentResponse
from app.services.audit import log_audit
from app.services.contract_docx import download_filename, save_contract_docx
from app.services.executor import build_document_variables

TEMPLATE_MAP = {
    DocumentType.CONTRACT: "contract_template.docx",
    DocumentType.INVOICE: "invoice_template.docx",
    DocumentType.ACT: "act_template.docx",
}


class DocumentService:
    def __init__(self, repo: DocumentRepository, contract_repo: ContractRepository):
        self.repo = repo
        self.contract_repo = contract_repo
        Path(settings.UPLOAD_DIR).mkdir(parents=True, exist_ok=True)
        get_templates_dir().mkdir(parents=True, exist_ok=True)

    def _can_edit(self, user: User, contract) -> bool:
        if not contract or not contract.client:
            return False
        return can_edit_contract(user, contract.client.owner_id, contract.responsible_manager_id)

    def _to_response(self, doc: Document) -> DocumentResponse:
        return DocumentResponse.model_validate(doc)

    async def list_documents(self, contract_id: int | None = None) -> list[DocumentResponse]:
        docs = await self.repo.list_all(contract_id=contract_id)
        return [self._to_response(d) for d in docs]

    async def get_document(self, doc_id: int) -> Document:
        doc = await self.repo.get_by_id(doc_id)
        if not doc:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Document not found")
        return doc

    async def _upsert_singleton_document(
        self,
        user: User,
        contract,
        doc_type: DocumentType,
        output_path: Path,
        db,
    ) -> Document:
        """Один документ данного типа на договор — повторная генерация обновляет файл."""
        existing = await self.repo.get_latest(contract.id, doc_type)
        if existing:
            old_path = Path(existing.file_path) if existing.file_path else None
            if old_path and old_path != output_path and old_path.is_file():
                old_path.unlink(missing_ok=True)
            existing.file_path = str(output_path)
            existing.generated_by = user.id
            existing.generated_at = datetime.now(timezone.utc)
            await db.flush()
            return existing

        document = Document(
            contract_id=contract.id,
            client_id=contract.client_id,
            type=doc_type,
            file_path=str(output_path),
            generated_by=user.id,
        )
        return await self.repo.create(document)

    async def build_contract_docx_file(
        self, user: User, contract_id: int, db
    ) -> tuple[Path, str]:
        contract = await self.contract_repo.get_by_id(contract_id)
        if not contract:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Contract not found")

        filename = download_filename(contract, contract.client)
        existing = await self.repo.get_latest(contract_id, DocumentType.CONTRACT)
        if existing and existing.file_path:
            output_path = Path(existing.file_path)
        else:
            output_path = Path(settings.UPLOAD_DIR) / f"contract_{contract.id}.docx"
        try:
            save_contract_docx(contract, contract.client, output_path)
        except FileNotFoundError as e:
            raise HTTPException(
                status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
                detail=str(e),
            ) from e

        document = await self._upsert_singleton_document(
            user, contract, DocumentType.CONTRACT, output_path, db
        )
        await log_audit(db, user.id, "GENERATE", "document", document.id)
        return output_path, filename

    def _substitute_docx(self, template_path: Path, output_path: Path, variables: dict[str, str]) -> None:
        from docx import Document as DocxDocument

        doc = DocxDocument(str(template_path))
        for paragraph in doc.paragraphs:
            for key, value in variables.items():
                placeholder = "{{" + key + "}}"
                if placeholder in paragraph.text:
                    paragraph.text = paragraph.text.replace(placeholder, value)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for key, value in variables.items():
                        placeholder = "{{" + key + "}}"
                        if placeholder in cell.text:
                            cell.text = cell.text.replace(placeholder, value)
        doc.save(str(output_path))

    def _ensure_template(self, doc_type: DocumentType) -> Path:
        template_name = TEMPLATE_MAP.get(doc_type)
        if not template_name:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Cannot generate this document type")
        template_path = get_templates_dir() / template_name
        if not template_path.exists():
            from docx import Document as DocxDocument

            doc = DocxDocument()
            doc.add_heading("Документ", 0)
            doc.add_paragraph("{{full_company_name}}")
            doc.add_paragraph("ИНН {{inn}} / КПП {{kpp}} / ОГРН {{ogrn}}")
            doc.add_paragraph("Адрес: {{legal_address}}")
            doc.add_paragraph("Банк: {{bank_name}}, БИК {{bik}}")
            doc.add_paragraph("р/с {{settlement_account}}, к/с {{correspondent_account}}")
            doc.add_paragraph("Договор № {{contract_number}} от {{date}}")
            doc.add_paragraph("Сумма: {{amount}}")
            doc.add_paragraph("Руководитель: {{director_name}}")
            doc.save(str(template_path))
        return template_path

    async def generate(self, user: User, data: DocumentGenerateRequest, db) -> DocumentResponse:
        contract = await self.contract_repo.get_by_id(data.contract_id)
        if not contract:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Contract not found")
        if not self._can_edit(user, contract):
            raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="No permission")

        template_path = self._ensure_template(data.document_type)
        if data.document_type == DocumentType.CONTRACT:
            output_path = Path(settings.UPLOAD_DIR) / f"contract_{contract.id}.docx"
        else:
            output_path = Path(settings.UPLOAD_DIR) / f"{data.document_type.value}_{contract.id}_{uuid4().hex[:8]}.docx"

        variables = build_document_variables(contract, contract.client)
        self._substitute_docx(template_path, output_path, variables)

        if data.document_type in (DocumentType.CONTRACT, DocumentType.ACT):
            document = await self._upsert_singleton_document(
                user, contract, data.document_type, output_path, db
            )
        else:
            document = Document(
                contract_id=contract.id,
                client_id=contract.client_id,
                type=data.document_type,
                file_path=str(output_path),
                generated_by=user.id,
            )
            document = await self.repo.create(document)
        await log_audit(db, user.id, "GENERATE", "document", document.id)
        return self._to_response(document)

    async def upload(
        self,
        user: User,
        contract_id: int,
        doc_type: DocumentType,
        file: UploadFile,
        db,
    ) -> DocumentResponse:
        contract = await self.contract_repo.get_by_id(contract_id)
        if not contract:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Contract not found")
        if not self._can_edit(user, contract):
            raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="No permission")

        ext = Path(file.filename or "file").suffix
        filename = f"upload_{contract_id}_{uuid4().hex[:8]}{ext}"
        output_path = Path(settings.UPLOAD_DIR) / filename
        with open(output_path, "wb") as f:
            shutil.copyfileobj(file.file, f)

        document = Document(
            contract_id=contract.id,
            client_id=contract.client_id,
            type=doc_type,
            file_path=str(output_path),
            generated_by=user.id,
        )
        document = await self.repo.create(document)
        await log_audit(db, user.id, "UPLOAD", "document", document.id)
        return self._to_response(document)
