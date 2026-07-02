"""Подстановка {{ переменных }} в Excel без потери форматирования ячеек."""
from __future__ import annotations

import re
from decimal import Decimal
from pathlib import Path
from typing import Any

import openpyxl
import xlrd
import xlwt
from xlutils.copy import copy
from xlutils.filter import BaseWriter
from xlwt.Style import default_style

_PLACEHOLDER = re.compile(r"\{\{\s*(\w+)\s*\}\}")

# Поля, для которых нужно писать в объединённую область шаблона.
_XLS_MERGE_ANCHORS = {
    "invoice_date_header",
    "buyer_name",
    "buyer_address",
    "buyer_inn_kpp",
    "line_text",
    "total_items_text",
    "amount_words",
}


def _merged_bounds_xlrd(ws_read: xlrd.sheet.Sheet, r: int, c: int) -> tuple[int, int, int, int] | None:
    for rlo, rhi, clo, chi in ws_read.merged_cells:
        if rlo <= r < rhi and clo <= c < chi:
            return rlo, rhi - 1, clo, chi - 1
    return None


def _lookup(name: str, variables: dict[str, Any]) -> Any:
    if name not in variables:
        return f"{{{{ {name} }}}}"
    return variables[name]


def substitute_placeholders(text: str, variables: dict[str, Any]) -> Any:
    """Заменить плейсхолдеры в строке. Если ячейка — один плейсхолдер, вернуть исходный тип."""
    stripped = text.strip()
    single = _PLACEHOLDER.fullmatch(stripped)
    if single:
        return _lookup(single.group(1), variables)

    def repl(match: re.Match[str]) -> str:
        value = _lookup(match.group(1), variables)
        if value is None:
            return ""
        return str(value)

    if "{{" not in text:
        return text
    return _PLACEHOLDER.sub(repl, text)


def _build_xls_styles(rb: xlrd.Book) -> list[Any]:
    writer = BaseWriter()
    writer.workbook(rb, "template.xls")
    return writer.style_list


def _xls_cell_style(rb: xlrd.Book, styles: list[Any], ws_read: xlrd.sheet.Sheet, r: int, c: int):
    try:
        xf_index = ws_read.cell_xf_index(r, c)
    except IndexError:
        return default_style
    if xf_index is None or not rb.formatting_info:
        return default_style
    return styles[xf_index]


def _normalize_xls_value(value: Any) -> Any:
    if isinstance(value, Decimal):
        return float(value)
    return value


def write_xls_cell(
    ws_write: xlwt.Worksheet,
    ws_read: xlrd.sheet.Sheet,
    rb: xlrd.Book,
    styles: list[Any],
    r: int,
    c: int,
    value: Any,
    *,
    field_key: str | None = None,
) -> None:
    """Записать значение, сохранив XF из шаблона (шрифт, формат числа, перенос текста)."""
    value = _normalize_xls_value(value)
    bounds = _merged_bounds_xlrd(ws_read, r, c) if field_key in _XLS_MERGE_ANCHORS else None
    if bounds:
        r, c = bounds[0], bounds[2]
    style = _xls_cell_style(rb, styles, ws_read, r, c)

    if bounds:
        r1, r2, c1, c2 = bounds
        for rr in range(r1, r2 + 1):
            for cc in range(c1, c2 + 1):
                ws_write.write(rr, cc, "", style)
        ws_write.write_merge(r1, r2, c1, c2, value, style)
        return

    ws_write.write(r, c, value, style)


def fill_xlsx_workbook(wb: openpyxl.Workbook, variables: dict[str, Any]) -> None:
    """Меняет только .value — шрифт, границы, ширина колонок и высота строк сохраняются."""
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if value is None or not isinstance(value, str) or "{{" not in value:
                    continue
                cell.value = substitute_placeholders(value, variables)


def _xlsx_has_placeholders(template_path: Path) -> bool:
    wb = openpyxl.load_workbook(str(template_path), read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                for value in row:
                    if isinstance(value, str) and "{{" in value:
                        return True
        return False
    finally:
        wb.close()


def _xls_has_placeholders(template_path: Path) -> bool:
    rb = xlrd.open_workbook(str(template_path))
    try:
        ws = rb.sheet_by_index(0)
        for r in range(ws.nrows):
            for c in range(ws.ncols):
                val = ws.cell_value(r, c)
                if isinstance(val, str) and "{{" in val:
                    return True
        return False
    finally:
        rb.release_resources()


# Координаты legacy-шаблона акта (1-based, openpyxl) — если в xlsx нет {{ }}
_LEGACY_ACT_ROWS = {
    "header": (1, 1),
    "preamble": (2, 1),
    "line_text": (6, 1),
    "line_unit": (6, 6),
    "line_qty": (6, 9),
    "line_price": (6, 11),
    "line_sum": (6, 13),
    "total_sum": (7, 13),
    "amount_words": (8, 1),
    "done_text": (10, 1),
    "executor_block": (12, 2),
    "customer_block": (12, 7),
    "sign_executor_label": (13, 2),
    "sign_executor_name": (13, 4),
    "sign_customer_label": (13, 8),
    "sign_customer_name": (13, 12),
}


def _fill_xlsx_legacy(template_path: Path, output_path: Path, variables: dict[str, Any]) -> None:
    wb = openpyxl.load_workbook(str(template_path))
    ws = wb.active
    mapping = {
        "header": variables.get("act_header", ""),
        "preamble": variables.get("act_preamble", ""),
        "line_text": variables.get("line_text", ""),
        "line_unit": variables.get("line_unit", ""),
        "line_qty": variables.get("line_qty", ""),
        "line_price": variables.get("line_price", ""),
        "line_sum": variables.get("line_sum", ""),
        "total_sum": variables.get("total_sum", ""),
        "amount_words": variables.get("amount_words", ""),
        "done_text": variables.get("done_text", ""),
        "executor_block": variables.get("executor_block", ""),
        "customer_block": variables.get("customer_block", ""),
        "sign_executor_label": variables.get("sign_executor_label", ""),
        "sign_executor_name": variables.get("sign_executor_name", ""),
        "sign_customer_label": variables.get("sign_customer_label", ""),
        "sign_customer_name": variables.get("sign_customer_name", ""),
    }
    for key, (row, col) in _LEGACY_ACT_ROWS.items():
        ws.cell(row, col).value = mapping[key]
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(output_path))


# Координаты legacy-шаблона счёта (0-based, xlwt) — под invoice_template.xls
_LEGACY_INVOICE_CELLS = {
    "invoice_date_header": (9, 0),
    "buyer_name": (12, 2),
    "buyer_address": (13, 2),
    "buyer_inn_kpp": (14, 2),
    "line_text": (17, 1),
    "line_qty": (17, 8),
    "line_price": (17, 9),
    "line_sum": (17, 13),
    "line_vat_pct": (17, 15),
    "line_vat": (17, 16),
    "line_total": (17, 18),
    "total_label": (18, 1),
    "total_qty": (18, 8),
    "total_price_dash": (18, 9),
    "total_sum": (18, 13),
    "total_vat_pct": (18, 15),
    "total_vat": (18, 16),
    "total_amount": (18, 18),
    "total_items_text": (20, 0),
    "amount_words": (21, 0),
}


def _fill_xls_legacy(template_path: Path, output_path: Path, variables: dict[str, Any]) -> None:
    rb = xlrd.open_workbook(str(template_path), formatting_info=True)
    wb = copy(rb)
    ws = wb.get_sheet(0)
    ws_read = rb.sheet_by_index(0)
    styles = _build_xls_styles(rb)
    for key, (row, col) in _LEGACY_INVOICE_CELLS.items():
        write_xls_cell(
            ws,
            ws_read,
            rb,
            styles,
            row,
            col,
            variables.get(key, ""),
            field_key=key,
        )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(output_path))


def fill_xlsx_template(template_path: Path, output_path: Path, variables: dict[str, Any]) -> None:
    if _xlsx_has_placeholders(template_path):
        wb = openpyxl.load_workbook(str(template_path))
        fill_xlsx_workbook(wb, variables)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(output_path))
    else:
        _fill_xlsx_legacy(template_path, output_path, variables)


def fill_xls_template(template_path: Path, output_path: Path, variables: dict[str, Any]) -> None:
    if _xls_has_placeholders(template_path):
        rb = xlrd.open_workbook(str(template_path), formatting_info=True)
        wb = copy(rb)
        ws_write = wb.get_sheet(0)
        ws_read = rb.sheet_by_index(0)
        styles = _build_xls_styles(rb)

        for r in range(ws_read.nrows):
            for c in range(ws_read.ncols):
                raw = ws_read.cell_value(r, c)
                if not isinstance(raw, str) or "{{" not in raw:
                    continue
                new_val = substitute_placeholders(raw, variables)
                write_xls_cell(ws_write, ws_read, rb, styles, r, c, new_val)

        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(output_path))
    else:
        _fill_xls_legacy(template_path, output_path, variables)

